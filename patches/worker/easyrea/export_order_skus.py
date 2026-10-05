#!/usr/bin/env python3
"""Read-only procurement demand from all outstanding Shopify order lines."""
from collections import defaultdict
from datetime import datetime
import csv
import os
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import runtime
from import_products import Shopify

VERSION = 1
SCOPES = 'query ProcurementScopes { currentAppInstallation { accessScopes { handle } } }'
LINE_FIELDS = '''id sku name vendor quantity currentQuantity unfulfilledQuantity requiresShipping isGiftCard
variant { id sku }'''
ORDERS = '''query ProcurementOrders($cursor: String, $query: String!) {
  orders(first: 20, after: $cursor, query: $query, sortKey: ID) {
    pageInfo { hasNextPage endCursor }
    nodes { id name createdAt cancelledAt test displayFinancialStatus displayFulfillmentStatus
      lineItems(first: 25) { pageInfo { hasNextPage endCursor } nodes { ''' + LINE_FIELDS + ''' } }
    }
  }
}'''
LINES = '''query ProcurementLines($id: ID!, $cursor: String) {
  order(id: $id) { lineItems(first: 100, after: $cursor) {
    pageInfo { hasNextPage endCursor } nodes { ''' + LINE_FIELDS + ''' }
  } }
}'''
# Include on-hold, scheduled and partially fulfilled orders, including closed orders.
SEARCH = '-status:cancelled -fulfillment_status:fulfilled'
SUMMARY_COLUMNS = ['SKU', 'Брой за изпращане', 'Бранд', 'Продукт', 'Брой поръчки', 'Най-стара поръчка', 'Поръчки', 'Бележка']
DETAIL_COLUMNS = ['Поръчка', 'Дата', 'SKU', 'Брой за изпращане', 'Бранд', 'Продукт', 'Плащане', 'Изпълнение', 'Бележка', 'Order ID', 'Line item ID', 'Variant ID']


def next_cursor(connection, previous, seen):
    if not isinstance(connection, dict) or not isinstance(connection.get('nodes'), list):
        raise RuntimeError('Непълна Shopify страница; отчетът не е създаден.')
    info = connection.get('pageInfo')
    if not isinstance(info, dict) or not isinstance(info.get('hasNextPage'), bool):
        raise RuntimeError('Липсва информация за Shopify пагинацията.')
    if not info['hasNextPage']:
        return None
    cursor = info.get('endCursor')
    if not cursor or cursor == previous or cursor in seen or not connection['nodes']:
        raise RuntimeError('Непълна Shopify пагинация; отчетът не е създаден.')
    seen.add(cursor)
    return cursor


def collect_orders(shop):
    cursor = None
    cursors, ids = set(), set()
    out = []
    while True:
        conn = shop.call(ORDERS, {'cursor': cursor, 'query': SEARCH}).get('orders')
        following = next_cursor(conn, cursor, cursors)
        for order in conn['nodes']:
            if order['id'] in ids:
                raise RuntimeError('Поръчка се повтаря между страници; пуснете отчета отново.')
            ids.add(order['id'])
            if order['cancelledAt'] or order['test'] or order['displayFinancialStatus'] in ('REFUNDED', 'VOIDED'):
                continue
            lines = order['lineItems']
            line_cursor = next_cursor(lines, None, set())
            line_cursors = {line_cursor} if line_cursor else set()
            all_lines = list(lines['nodes'])
            while line_cursor:
                data = shop.call(LINES, {'id': order['id'], 'cursor': line_cursor}).get('order')
                if not data:
                    raise RuntimeError('Поръчката е променена или вече не е достъпна; пуснете отчета отново.')
                page = data['lineItems']
                following_line = next_cursor(page, line_cursor, line_cursors)
                all_lines.extend(page['nodes'])
                line_cursor = following_line
            if len({line['id'] for line in all_lines}) != len(all_lines):
                raise RuntimeError('Повторени редове в поръчка; пуснете отчета отново.')
            out.append({**order, 'items': all_lines})
        print(f'[shopify] Проверени поръчки: {len(ids)}', flush=True)
        if not following:
            return out, len(ids)
        cursor = following


def date_label(value):
    return datetime.fromisoformat(value.replace('Z', '+00:00')).astimezone(ZoneInfo('Europe/Sofia')).strftime('%Y-%m-%d %H:%M')


def build_rows(orders):
    groups = defaultdict(list)
    details, review = [], []
    for order in orders:
        if order['cancelledAt'] or order['test'] or order['displayFinancialStatus'] in ('REFUNDED', 'VOIDED'):
            continue
        for line in order['items']:
            if not line['requiresShipping'] or line['isGiftCard']:
                continue
            # Shopify's outstanding quantity is authoritative; never use original ordered units.
            # Cap at the current quantity, which excludes refunded and removed units.
            qty = max(0, min(line['unfulfilledQuantity'], line['currentQuantity']))
            if not qty:
                continue
            variant = line.get('variant') or {}
            ordered_sku = (line.get('sku') or '').strip()
            current_sku = (variant.get('sku') or '').strip()
            code = ordered_sku or current_sku
            notes = []
            if not code: notes.append('Липсва SKU; нужна е ръчна проверка')
            elif not ordered_sku: notes.append('SKU е взет от текущия вариант')
            if not variant: notes.append('Продуктът/вариантът е изтрит; използвани са данните от поръчката')
            if ordered_sku and current_sku and ordered_sku != current_sku:
                notes.append('SKU е променен: в поръчката ' + ordered_sku + ', текущ ' + current_sku)
            if not line.get('vendor'): notes.append('Липсва бранд')
            row = {'Поръчка': order['name'], 'Дата': date_label(order['createdAt']), 'SKU': code,
                   'Брой за изпращане': qty, 'Бранд': (line.get('vendor') or '').strip(), 'Продукт': line['name'],
                   'Плащане': order['displayFinancialStatus'], 'Изпълнение': order['displayFulfillmentStatus'],
                   'Бележка': '; '.join(notes), 'Order ID': order['id'], 'Line item ID': line['id'], 'Variant ID': variant.get('id', '')}
            details.append(row)
            if notes: review.append(dict(row))
            if code: groups[code].append(row)
    summary = []
    for code, rows in groups.items():
        vendors = sorted({r['Бранд'] for r in rows if r['Бранд']})
        variants = {r['Variant ID'] for r in rows if r['Variant ID']}
        notes = sorted({r['Бележка'] for r in rows if r['Бележка']})
        if len(vendors) > 1 or len(variants) > 1:
            conflict = 'SKU се среща при различни брандове/варианти; проверете преди заявяване'
            notes.append(conflict)
            for row in rows:
                review.append({**row, 'Бележка': '; '.join(filter(None, [row['Бележка'], conflict]))})
        names = sorted({r['Поръчка'] for r in rows})
        summary.append({'SKU': code, 'Брой за изпращане': sum(r['Брой за изпращане'] for r in rows),
                        'Бранд': ' | '.join(vendors), 'Продукт': ' | '.join(sorted({r['Продукт'] for r in rows})),
                        'Брой поръчки': len({r['Order ID'] for r in rows}), 'Най-стара поръчка': min(r['Дата'] for r in rows),
                        'Поръчки': ', '.join(names), 'Бележка': '; '.join(notes)})
    summary.sort(key=lambda r: (r['Бранд'].casefold(), r['SKU']))
    details.sort(key=lambda r: (r['Дата'], r['Поръчка'], r['SKU']))
    return summary, details, review


def add_sheet(book, title, columns, rows):
    sheet = book.create_sheet(title)
    sheet.append(columns)
    for row in rows:
        sheet.append([row.get(c, '') for c in columns])
    # Keep SKU, leading zeros and all source text as literal text, not spreadsheet formulas.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str): cell.data_type = 's'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='263F33')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for i, name in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(i)].width = 45 if name in ('Продукт', 'Бележка', 'Поръчки') else 23
    return sheet


def write_reports(summary, details, review, directory, scanned):
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(ZoneInfo('Europe/Sofia')).strftime('%Y%m%d_%H%M%S')
    base = directory / ('unshipped_order_skus_' + stamp)
    book = Workbook(); book.remove(book.active)
    add_sheet(book, 'SKU обобщение', SUMMARY_COLUMNS, summary)
    add_sheet(book, 'По поръчки', DETAIL_COLUMNS, details)
    add_sheet(book, 'За проверка', DETAIL_COLUMNS, review)
    metadata = [
        ('Създаден на', datetime.now(ZoneInfo('Europe/Sofia')).isoformat(timespec='seconds')),
        ('Обхват', 'Всички достъпни неизпратени и частично изпратени поръчки, без ограничение по дата'),
        ('Проверени поръчки', scanned),
        ('Поръчки с оставащи физически продукти', len({r['Order ID'] for r in details})),
        ('Различни SKU', len(summary)),
        ('Бройки със SKU', sum(r['Брой за изпращане'] for r in summary)),
        ('Бройки без SKU', sum(r['Брой за изпращане'] for r in details if not r['SKU'])),
        ('Изключени', 'Анулирани, тестови, изцяло възстановени и voided поръчки; изпратени редове, услуги и подаръчни карти'),
        ('Плащане', 'Включени са и неплатени поръчки / наложен платеж; статусът е в лист По поръчки'),
        ('Количество', 'Оставаща клиентска нужда. Не са приспаднати складови наличности и вече заявена стока.'),
        ('Бранд', 'Shopify Vendor; брандът не е автоматично съпоставен с търговски доставчик'),
        ('Проверка', 'SKU с бележка и лист За проверка се преглеждат преди заявяване'),
        ('Права', 'Проверени read_orders/write_orders и read_all_orders; отчетът не променя Shopify'),
    ]
    add_sheet(book, 'Информация', ['Поле', 'Стойност'], [{'Поле': k, 'Стойност': v} for k, v in metadata])
    book.save(base.with_suffix('.xlsx'))
    # SKU and quantity first for quick copy, with review notes retained.
    with base.with_suffix('.csv').open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f); writer.writerow(['SKU', 'Количество', 'Бранд', 'Бележка'])
        for row in summary:
            code = row['SKU']
            writer.writerow(["'" + code if code.startswith(('=', '+', '-', '@', '\t', '\r')) else code, row['Брой за изпращане'], row['Бранд'], row['Бележка']])
    return base.with_suffix('.xlsx'), base.with_suffix('.csv')


def main():
    config = runtime.require_env('SHOPIFY_STORE', 'SHOPIFY_CLIENT_ID', 'SHOPIFY_CLIENT_SECRET')
    print('=== SKU от неизпратени поръчки — само четене ===', flush=True)
    shop = Shopify(config)
    scopes = {s['handle'] for s in shop.call(SCOPES)['currentAppInstallation']['accessScopes']}
    if not scopes.intersection({'read_orders', 'write_orders'}):
        raise RuntimeError('Липсва read_orders в Shopify приложението на Hub. Разрешете четене на поръчки.')
    if 'read_all_orders' not in scopes:
        raise RuntimeError('Липсва read_all_orders в Shopify приложението на Hub. Без него Shopify връща само последните 60 дни. Разрешете достъп до всички поръчки, за да получите пълна заявка.')
    orders, scanned = collect_orders(shop)
    summary, details, review = build_rows(orders)
    outputs = write_reports(summary, details, review, os.environ.get('HUB_REPORT_DIR', '.'), scanned)
    print(f'[резултат] {len(summary)} SKU | {sum(r["Брой за изпращане"] for r in summary)} броя със SKU | {len({r["Order ID"] for r in details})} поръчки | {len(review)} реда за проверка', flush=True)
    print('Количествата са клиентска нужда; наличностите и вече заявената стока не са приспаднати.', flush=True)
    for output in outputs: print(f'Отчет записан: {output}', flush=True)


if __name__ == '__main__': main()
