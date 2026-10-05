import sys
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'easyrea'))
from export_order_skus import build_rows, collect_orders, write_reports, main, ORDERS, LINES
from openpyxl import load_workbook


def line(id='l1', sku='00123', qty=3, outstanding=2, current=3, vendor='atmosphera', variant='v1'):
    return {'id': id, 'sku': sku, 'name': 'Стол', 'vendor': vendor, 'quantity': qty,
            'currentQuantity': current, 'unfulfilledQuantity': outstanding, 'requiresShipping': True,
            'isGiftCard': False, 'variant': {'id': variant, 'sku': sku} if variant else None}


def order(id='o1', items=None, **changes):
    return {'id': id, 'name': '#'+id, 'createdAt': '2026-01-02T10:00:00Z', 'cancelledAt': None, 'test': False,
            'displayFinancialStatus': 'PENDING', 'displayFulfillmentStatus': 'PARTIALLY_FULFILLED',
            'items': items if items is not None else [line()], **changes}


def connection(nodes, cursor=None):
    return {'nodes': nodes, 'pageInfo': {'hasNextPage': cursor is not None, 'endCursor': cursor}}


class OrderSkus(unittest.TestCase):
    def test_same_sku_sums_outstanding_not_original_and_keeps_pending_cod(self):
        summary, detail, review = build_rows([order(), order('o2', [line('l2', qty=8, outstanding=1, current=8)])])
        self.assertEqual(summary[0]['SKU'], '00123')
        self.assertEqual(summary[0]['Брой за изпращане'], 3)
        self.assertEqual(summary[0]['Брой поръчки'], 2)
        self.assertEqual(review, [])
    def test_cancelled_test_voided_refunded_shipped_and_nonphysical_excluded(self):
        items = [line(outstanding=0), {**line('l2'), 'isGiftCard': True}, {**line('l3'), 'requiresShipping': False}, line('l4', current=0)]
        summary, details, _ = build_rows([order(items=items), order('o2', cancelledAt='date'), order('o3', test=True), order('o4', displayFinancialStatus='REFUNDED'), order('o5', displayFinancialStatus='VOIDED')])
        self.assertEqual((summary, details), ([], []))
    def test_deleted_variant_sku_and_missing_sku_are_not_lost(self):
        summary, details, review = build_rows([order(items=[line(variant=None), line('l2', sku='', variant=None)])])
        self.assertEqual(len(summary), 1); self.assertEqual(len(details), 2); self.assertEqual(len(review), 2)
        self.assertIn('изтрит', review[0]['Бележка'])
    def test_changed_sku_is_flagged_and_historical_code_kept(self):
        summary, _, review = build_rows([order(items=[{**line(), 'variant': {'id': 'v1', 'sku': 'NEW'}}])])
        self.assertEqual(summary[0]['SKU'], '00123'); self.assertIn('променен', review[0]['Бележка'])
    def test_same_sku_in_multiple_variants_requires_review(self):
        summary, _, review = build_rows([order(), order('o2', [line('l2', variant='v2', vendor='5five')])])
        self.assertEqual(summary[0]['Брой за изпращане'], 4)
        self.assertIn('различни', summary[0]['Бележка']); self.assertEqual(len(review), 2)
    def test_both_orders_and_nested_line_items_paginate(self):
        o1 = order(); o1['lineItems'] = connection([line()], 'line-next')
        o2 = order('o2'); o2['lineItems'] = connection([line('l3')])
        shop = Mock(); shop.call.side_effect = [
            {'orders': connection([o1], 'order-next')},
            {'order': {'lineItems': connection([line('l2')])}},
            {'orders': connection([o2])}]
        orders, scanned = collect_orders(shop)
        self.assertEqual(scanned, 2); self.assertEqual(len(orders[0]['items']), 2)
        self.assertEqual(shop.call.call_args_list[1].args, (LINES, {'id': 'o1', 'cursor': 'line-next'}))
        self.assertEqual(shop.call.call_args_list[2].args[1]['cursor'], 'order-next')
    def test_cursor_cycle_fails_instead_of_partial_report(self):
        shop = Mock(); shop.call.side_effect = [
            {'orders': connection([], 'same')},
        ]
        with self.assertRaises(RuntimeError): collect_orders(shop)
    def test_scope_gap_blocks_incomplete_60_day_export(self):
        shop = Mock(); shop.call.return_value = {'currentAppInstallation': {'accessScopes': [{'handle': 'read_orders'}]}}
        with patch('export_order_skus.runtime.require_env', return_value={}), patch('export_order_skus.Shopify', return_value=shop), patch('export_order_skus.collect_orders') as collect:
            with self.assertRaisesRegex(RuntimeError, 'read_all_orders'): main()
            collect.assert_not_called()
    def test_xlsx_preserves_zero_sku_and_literal_formula_and_has_all_sheets(self):
        summary, details, review = build_rows([order(items=[line(sku='00123'), {**line('l2', sku='=1+2'), 'name': '=HYPERLINK("bad")'}])])
        with tempfile.TemporaryDirectory() as d:
            xlsx, csv = write_reports(summary, details, review, d, 1)
            book = load_workbook(xlsx)
            self.assertEqual(book.sheetnames, ['SKU обобщение', 'По поръчки', 'За проверка', 'Информация'])
            cells = list(book['SKU обобщение'].values)
            self.assertEqual(cells[1][0], '00123')
            self.assertEqual(book['SKU обобщение']['A3'].data_type, 's')
            self.assertTrue(book['SKU обобщение'].auto_filter.ref)
            self.assertIn('00123', csv.read_text(encoding='utf-8-sig'))
    def test_empty_report_still_has_headers(self):
        with tempfile.TemporaryDirectory() as d:
            xlsx, _ = write_reports([], [], [], d, 0)
            self.assertEqual(load_workbook(xlsx)['SKU обобщение'].max_row, 1)

if __name__ == '__main__': unittest.main()
